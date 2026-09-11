"""
Extract unique program names from Excel files, categorize as real vs garbage,
and output a clean list.
"""
import json
import re
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import load_workbook

# Patterns that indicate a page is NOT a real degree program
GARBAGE_PATTERNS = [
    # Navigation / generic pages
    r"^(Research|Programmes|Courses|Team|Facilities|News|Events|About|Home|Contact)$",
    r"^(Undergraduate courses|Postgraduate programmes|Postgraduate Programmes)$",
    r"^(Postgraduate Research|Current Students|Scholarships|Funding your programme)$",
    r"^(Image Gallery|Video Gallery|Photo Gallery|Gallery)$",
    r"^(Staff|Faculty|People|Our People|Our Team|Our Staff)$",
    r"^(Apply|How to Apply|Application|Admissions)$",
    r"^(Major Map)$",
    
    # Policy / admin pages
    r"(?i)(policy|policies|procedure|handbook|handbook|governance|guideline)",
    r"(?i)(privacy|cookie|disclaimer|terms of use|terms and conditions)",
    
    # Country / location pages (UL international)
    r"^(Canada|Japan|China|India|USA|UK|Germany|France|Brazil|Mexico|Nigeria|Korea|Malaysia)$",
    r"^(Saudi Arabia|Pakistan|Vietnam|Indonesia|Thailand|Bangladesh|Sri Lanka|Nepal)$",
    r"^(Hong Kong|Singapore|Taiwan|Philippines|Turkey|Iran|Iraq|Egypt|Kenya|Ghana)$",
    
    # Staff profile pages
    r"(?i)^(Dr\.?|Prof\.?|Professor) ",
    
    # Frankfurt navigation duplicates
    r"(?i)^(Perspectives|Is this degree program right for me|Advisory Services and Contacts)$",
    r"(?i)^(Is the Degree Programme Right for Me|Advising and Contact)$",
    r"(?i)^(Application and Admission|Application and Admission Requirements)$",
    
    # Generic descriptions
    r"(?i)(training course catalogue|full degree|graduate profiles?$)",
    r"(?i)^(Overview|Introduction|Welcome|FAQ|Help)$",
]

# Patterns that strongly suggest it IS a real degree program
PROGRAM_PATTERNS = [
    r",\s*(B[AS]|BS|BA|BEng|BCom|BDes|BFA|BMusc?)$",  # Bachelor's
    r",\s*(M[AS]|MS|MA|MSc|MBA|MEng|MFA|MPhil|LLM)$",  # Master's
    r",\s*(PhD|DBA|EdD|DrPH|JD|MD)$",  # Doctoral
    r"\b(Bachelor|Master|Doctor|Diploma|Certificate)\b",
    r"\b(BSc|BA|MA|MSc|PhD|BEng|MEng|LLB|LLM|MPhil|BComm?)\b",
    r"(?i)\b(degree|program|programme|major|minor)\b",
    r"- (B[AS]|M[AS]|PhD|BSc|MSc|BA|MA|BEng|MEng)\b",
]


def is_garbage(name: str) -> bool:
    """Check if a program name matches garbage patterns."""
    if not name or not name.strip():
        return True
    for pattern in GARBAGE_PATTERNS:
        if re.search(pattern, name.strip()):
            return True
    return False


def is_likely_program(name: str) -> bool:
    """Check if a program name looks like a real degree program."""
    for pattern in PROGRAM_PATTERNS:
        if re.search(pattern, name.strip()):
            return True
    return False


def extract_programs(filepath: str) -> dict:
    """Extract and categorize programs from an Excel file."""
    wb = load_workbook(filepath, read_only=True, data_only=True)
    
    data_sheet = None
    for name in wb.sheetnames:
        if name != "Summary":
            data_sheet = wb[name]
            break
    
    rows = list(data_sheet.iter_rows(values_only=True))
    headers = [str(h).strip().lower().replace(" ", "_") if h else "" for h in rows[0]]
    
    pid_col = headers.index("program_id") if "program_id" in headers else 0
    pname_col = headers.index("program_name") if "program_name" in headers else 1
    field_col = headers.index("field") if "field" in headers else 4
    
    # Collect unique programs with fact counts
    programs = {}  # pid -> {name, fact_count, has_programme_name_field}
    
    for row in rows[1:]:
        if not row or not row[pid_col]:
            continue
        pid = str(row[pid_col]).strip()
        pname = str(row[pname_col]).strip() if row[pname_col] else ""
        field = str(row[field_col]).strip() if row[field_col] else ""
        
        if pid not in programs:
            programs[pid] = {"name": pname, "fact_count": 0, "has_programme_name": False}
        programs[pid]["fact_count"] += 1
        if field == "programme_name":
            programs[pid]["has_programme_name"] = True
    
    wb.close()
    
    # Categorize
    real_programs = []
    garbage = []
    uncertain = []
    
    for pid, info in sorted(programs.items()):
        name = info["name"]
        facts = info["fact_count"]
        
        if is_garbage(name):
            garbage.append((pid, name, facts))
        elif is_likely_program(name) and facts >= 4:
            real_programs.append((pid, name, facts))
        elif facts >= 10 and info["has_programme_name"]:
            real_programs.append((pid, name, facts))
        elif facts <= 3:
            garbage.append((pid, name, facts))
        else:
            uncertain.append((pid, name, facts))
    
    return {
        "real": real_programs,
        "garbage": garbage,
        "uncertain": uncertain,
        "total_in_excel": len(programs),
    }


def print_results(uni_name: str, results: dict):
    print(f"\n{'='*70}")
    print(f"  {uni_name}")
    print(f"{'='*70}")
    print(f"  Total in Excel:     {results['total_in_excel']}")
    print(f"  Real Programs:      {len(results['real'])}")
    print(f"  Garbage:            {len(results['garbage'])}")
    print(f"  Uncertain:          {len(results['uncertain'])}")
    
    print(f"\n  REAL PROGRAMS ({len(results['real'])}):")
    print(f"  {'─'*60}")
    for pid, name, facts in sorted(results["real"], key=lambda x: x[1]):
        print(f"    [{pid}] ({facts:>3} facts) {name[:90]}")
    
    if results["uncertain"]:
        print(f"\n  UNCERTAIN ({len(results['uncertain'])}) - Need manual review:")
        print(f"  {'─'*60}")
        for pid, name, facts in sorted(results["uncertain"], key=lambda x: -x[2])[:30]:
            print(f"    [{pid}] ({facts:>3} facts) {name[:90]}")
    
    print(f"\n  GARBAGE SAMPLE ({min(20, len(results['garbage']))} of {len(results['garbage'])}):")
    print(f"  {'─'*60}")
    for pid, name, facts in sorted(results["garbage"], key=lambda x: -x[2])[:20]:
        print(f"    [{pid}] ({facts:>3} facts) {name[:90]}")


if __name__ == "__main__":
    files = {
        "ASU": "asu_extraction_data.xlsx",
        "LIMERICK (UL)": "limerick_extraction_data.xlsx",
        "FRANKFURT": "frankfurt_extraction_data.xlsx",
    }
    
    all_results = {}
    for uni, filepath in files.items():
        if Path(filepath).exists():
            results = extract_programs(filepath)
            all_results[uni] = results
            print_results(uni, results)
        else:
            print(f"\n  [SKIP] {filepath} not found")
    
    # Final summary
    print(f"\n{'='*70}")
    print(f"  FINAL SUMMARY")
    print(f"{'='*70}")
    for uni, r in all_results.items():
        print(f"  {uni:20s} | In Excel: {r['total_in_excel']:>5} | Real: {len(r['real']):>5} | Garbage: {len(r['garbage']):>5} | Uncertain: {len(r['uncertain']):>5}")
