"""
Analyze the Excel extraction files to understand program counts and data quality.
"""
import json
from pathlib import Path
from collections import Counter, defaultdict

from openpyxl import load_workbook


def analyze_excel(filepath: str):
    print(f"\n{'='*70}")
    print(f"  Analyzing: {filepath}")
    print(f"{'='*70}")
    
    wb = load_workbook(filepath, read_only=True, data_only=True)
    
    # Find the data sheet (not Summary)
    data_sheet = None
    for name in wb.sheetnames:
        if "All Facts" in name or name != "Summary":
            data_sheet = wb[name]
            break
    
    if not data_sheet:
        print("  [ERROR] No data sheet found")
        return
    
    print(f"  Sheet: {data_sheet.title}")
    
    # Read headers
    rows = list(data_sheet.iter_rows(values_only=True))
    if not rows:
        print("  [ERROR] Empty sheet")
        return
    
    headers = [str(h).strip() if h else "" for h in rows[0]]
    print(f"  Headers: {headers}")
    print(f"  Total rows (excl header): {len(rows) - 1}")
    
    # Find column indices
    col_map = {}
    for i, h in enumerate(headers):
        col_map[h.lower().replace(" ", "_")] = i
    
    pid_col = col_map.get("program_id", 0)
    pname_col = col_map.get("program_name", 1)
    cat_col = col_map.get("category", 2)
    subcat_col = col_map.get("subcategory", 3)
    field_col = col_map.get("field", 4)
    value_col = col_map.get("value", 5)
    url_col = col_map.get("source_url", 6)
    
    # Collect stats
    programs = {}  # program_id -> program_name
    subcategory_counts = Counter()
    field_counts = Counter()
    facts_per_program = Counter()
    programs_with_urls = set()
    total_with_url = 0
    total_without_url = 0
    
    # Sample some program names to check for non-program entries
    sample_names = []
    
    for row in rows[1:]:
        if not row or not row[pid_col]:
            continue
        
        pid = str(row[pid_col]).strip()
        pname = str(row[pname_col]).strip() if row[pname_col] else ""
        subcat = str(row[subcat_col]).strip() if row[subcat_col] else ""
        field = str(row[field_col]).strip() if row[field_col] else ""
        url = str(row[url_col]).strip() if len(row) > url_col and row[url_col] else ""
        
        programs[pid] = pname
        subcategory_counts[subcat] += 1
        field_counts[field] += 1
        facts_per_program[pid] += 1
        
        if url and url != "None" and url != "":
            total_with_url += 1
            programs_with_urls.add(pid)
        else:
            total_without_url += 1
    
    # ── Results ──────────────────────────────────────────────────
    total_programs = len(programs)
    total_facts = len(rows) - 1
    
    print(f"\n  PROGRAM STATS:")
    print(f"  ─────────────")
    print(f"  Unique programs: {total_programs}")
    print(f"  Total facts: {total_facts}")
    print(f"  Avg facts/program: {total_facts / max(total_programs, 1):.1f}")
    
    # Facts per program distribution
    fact_counts = list(facts_per_program.values())
    fact_counts.sort()
    
    zero_fact_programs = sum(1 for c in fact_counts if c == 0)
    low_fact_programs = sum(1 for c in fact_counts if 0 < c <= 3)
    medium_fact_programs = sum(1 for c in fact_counts if 3 < c <= 20)
    high_fact_programs = sum(1 for c in fact_counts if c > 20)
    
    print(f"\n  FACTS PER PROGRAM DISTRIBUTION:")
    print(f"  ───────────────────────────────")
    print(f"  0 facts:      {zero_fact_programs} programs")
    print(f"  1-3 facts:    {low_fact_programs} programs  ← SUSPICIOUS (likely not real programs)")
    print(f"  4-20 facts:   {medium_fact_programs} programs")
    print(f"  20+ facts:    {high_fact_programs} programs")
    print(f"  Min: {min(fact_counts)}, Max: {max(fact_counts)}, Median: {fact_counts[len(fact_counts)//2]}")
    
    # Source URL coverage
    print(f"\n  SOURCE URL COVERAGE:")
    print(f"  ────────────────────")
    print(f"  Facts with URL:    {total_with_url} ({100*total_with_url/max(total_facts,1):.1f}%)")
    print(f"  Facts without URL: {total_without_url} ({100*total_without_url/max(total_facts,1):.1f}%)")
    print(f"  Programs with at least 1 URL: {len(programs_with_urls)}/{total_programs}")
    
    # Subcategory breakdown
    print(f"\n  SUBCATEGORY BREAKDOWN:")
    print(f"  ──────────────────────")
    for subcat, count in subcategory_counts.most_common(15):
        print(f"    {subcat:30s} {count:>6,}")
    
    # Top fields
    print(f"\n  TOP FIELDS:")
    print(f"  ───────────")
    for field, count in field_counts.most_common(15):
        print(f"    {field:40s} {count:>6,}")
    
    # Sample programs with very few facts (likely garbage)
    print(f"\n  SAMPLE LOW-FACT PROGRAMS (1-2 facts — possibly not real programs):")
    print(f"  ──────────────────────────────────────────────────────────────────")
    count = 0
    for pid, fc in sorted(facts_per_program.items(), key=lambda x: x[1]):
        if fc <= 2 and count < 15:
            print(f"    [{pid}] ({fc} facts) {programs[pid][:80]}")
            count += 1
    
    # Sample programs with many facts (likely legit)
    print(f"\n  SAMPLE HIGH-FACT PROGRAMS (legit programs):")
    print(f"  ────────────────────────────────────────────")
    count = 0
    for pid, fc in sorted(facts_per_program.items(), key=lambda x: -x[1]):
        if count < 10:
            print(f"    [{pid}] ({fc} facts) {programs[pid][:80]}")
            count += 1
    
    # Check for duplicate program names
    name_counts = Counter(programs.values())
    duplicates = {name: count for name, count in name_counts.items() if count > 1}
    if duplicates:
        print(f"\n  DUPLICATE PROGRAM NAMES ({len(duplicates)} names appear multiple times):")
        print(f"  ─────────────────────────────────────────────────────────────")
        for name, count in sorted(duplicates.items(), key=lambda x: -x[1])[:15]:
            print(f"    ({count}x) {name[:80]}")
    
    wb.close()
    return {
        "total_programs": total_programs,
        "total_facts": total_facts,
        "low_fact_programs": low_fact_programs,
        "high_fact_programs": high_fact_programs,
    }


if __name__ == "__main__":
    files = [
        "asu_extraction_data.xlsx",
        "limerick_extraction_data.xlsx",
        "frankfurt_extraction_data.xlsx",
    ]
    
    for f in files:
        if Path(f).exists():
            analyze_excel(f)
        else:
            print(f"\n  [SKIP] {f} not found")
