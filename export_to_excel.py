"""
export_to_excel.py

Export extracted university programme data to an Excel file.

Each row represents a single extracted fact with columns:
  - Program ID
  - Program Name
  - Category (e.g. programme, university)
  - Subcategory (e.g. fees, identity, overview)
  - Field (e.g. tuition_fee, programme_name)
  - Value
  - Source URL
  - Source Type
  - Confidence

Usage:
    python export_to_excel.py --university asu
    python export_to_excel.py --university asu --output asu_data.xlsx
"""

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── University path mapping ──────────────────────────────────────
UNIVERSITY_PATHS = {
    "asu": Path("data/united_states/asu"),
    "frankfurt": Path("data/germany/uni-frankfurt"),
    "limerick": Path("data/ireland/ul"),
}

# ── JSON files to read per program (order matters for sheet logic) ──
FACT_FILES = [
    "program_data.json",
    "curriculum_data.json",
    "fees_data.json",
    "career_data.json",
    "contacts_data.json",
    "scholarships_data.json",
    "statistics_data.json",
    "student_life_data.json",
    "other_data.json",
]


def get_source_url_for_chunk(program_dir: Path, chunk_id: str) -> str:
    """Resolve a chunk_id to its actual source URL."""
    if not chunk_id:
        return ""

    try:
        if chunk_id.startswith("program_"):
            # Main program URL
            meta_path = program_dir / "metadata.json"
            if meta_path.exists():
                with open(meta_path, "r", encoding="utf-8") as f:
                    return json.load(f).get("url", "")
                    
        elif chunk_id.startswith("page_"):
            # Sub-page URL
            page_id = chunk_id.replace("page_", "")
            meta_path = program_dir / "pages" / page_id / "metadata.json"
            if meta_path.exists():
                with open(meta_path, "r", encoding="utf-8") as f:
                    return json.load(f).get("url", "")
                    
        elif chunk_id.startswith("fallback_"):
            # Fallback search URL
            # The chunk ID might be something like fallback_3f1d03b4ca
            # which directly matches the folder name
            meta_path = program_dir / "pages" / chunk_id / "metadata.json"
            if meta_path.exists():
                with open(meta_path, "r", encoding="utf-8") as f:
                    return json.load(f).get("url", "")
                    
        elif chunk_id.startswith("pdf_"):
            # PDF URL
            pdf_id = chunk_id.replace("pdf_", "")
            meta_path = program_dir / "pdf" / pdf_id / "extracted" / "document_data.json"
            if meta_path.exists():
                with open(meta_path, "r", encoding="utf-8") as f:
                    return json.load(f).get("source_url", "")
                    
    except Exception:
        pass
        
    return ""


def load_facts_from_file(path: Path, program_dir: Path) -> list[dict]:
    """Load facts from a single JSON file and return flat list of fact dicts."""
    if not path.exists():
        return []

    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return []

    facts = []
    program_id = data.get("program_id", "")
    program_name = data.get("program_name", "")

    # Collect from both university_facts and programme_facts
    for fact_list_key in ("university_facts", "programme_facts"):
        for fact in data.get(fact_list_key, []):
            source_url = fact.get("source_url", "")
            
            # If empty, try to resolve it from the chunk_id in metadata
            if not source_url:
                metadata = fact.get("metadata", {})
                chunk_id = metadata.get("chunk_id", "")
                if chunk_id:
                    source_url = get_source_url_for_chunk(program_dir, chunk_id)

            facts.append({
                "program_id": program_id,
                "program_name": program_name,
                "category": fact.get("category", ""),
                "subcategory": fact.get("subcategory", ""),
                "field": fact.get("field", ""),
                "value": str(fact.get("value", "")),
                "source_url": source_url,
                "source_type": fact.get("source_type", ""),
                "confidence": fact.get("confidence", ""),
            })

    return facts


def collect_all_facts(university_path: Path) -> list[dict]:
    """Iterate over all programs and collect every fact."""
    programs_dir = university_path / "programs"
    if not programs_dir.exists():
        print(f"  [ERROR] Programs directory not found: {programs_dir}")
        return []

    all_facts = []
    program_dirs = sorted(programs_dir.iterdir())

    for program_dir in program_dirs:
        if not program_dir.is_dir():
            continue

        final_dir = program_dir / "final"
        if not final_dir.exists():
            continue

        for fact_file in FACT_FILES:
            facts = load_facts_from_file(final_dir / fact_file, program_dir)
            all_facts.extend(facts)

    return all_facts


def write_excel(facts: list[dict], output_path: Path, university_name: str):
    """Write facts to a styled Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"{university_name.upper()} - All Facts"

    # ── Column headers ───────────────────────────────────────────
    headers = [
        "Program ID",
        "Program Name",
        "Category",
        "Subcategory",
        "Field",
        "Value",
        "Source URL",
        "Source Type",
        "Confidence",
    ]

    # ── Styling ──────────────────────────────────────────────────
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="8B1A1A", end_color="8B1A1A", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    even_fill = PatternFill(start_color="FFF2F2", end_color="FFF2F2", fill_type="solid")
    odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Write headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data rows
    data_font = Font(name="Calibri", size=10)
    wrap_alignment = Alignment(vertical="top", wrap_text=True)

    for row_idx, fact in enumerate(facts, start=2):
        fill = even_fill if row_idx % 2 == 0 else odd_fill

        values = [
            fact["program_id"],
            fact["program_name"],
            fact["category"],
            fact["subcategory"],
            fact["field"],
            fact["value"],
            fact["source_url"],
            fact["source_type"],
            fact["confidence"],
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = wrap_alignment

    # ── Column widths ────────────────────────────────────────────
    col_widths = [12, 45, 14, 18, 30, 60, 50, 14, 12]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(facts) + 1}"

    # ── Summary sheet ────────────────────────────────────────────
    ws_summary = wb.create_sheet(title="Summary", index=0)

    # Collect summary stats
    programs = set()
    categories = {}
    for fact in facts:
        programs.add(fact["program_id"])
        cat = fact["subcategory"] or fact["category"]
        categories[cat] = categories.get(cat, 0) + 1

    summary_data = [
        ("University", university_name.upper()),
        ("Total Programs", len(programs)),
        ("Total Facts", len(facts)),
        ("", ""),
        ("Category Breakdown", "Count"),
    ]

    for cat in sorted(categories.keys()):
        summary_data.append((cat, categories[cat]))

    title_font = Font(name="Calibri", bold=True, size=14, color="8B1A1A")
    label_font = Font(name="Calibri", bold=True, size=11)
    value_font = Font(name="Calibri", size=11)

    ws_summary.column_dimensions["A"].width = 35
    ws_summary.column_dimensions["B"].width = 20

    ws_summary.cell(row=1, column=1, value=f"{university_name.upper()} — Extraction Summary").font = title_font
    ws_summary.merge_cells("A1:B1")

    for row_idx, (label, value) in enumerate(summary_data, start=3):
        cell_a = ws_summary.cell(row=row_idx, column=1, value=label)
        cell_b = ws_summary.cell(row=row_idx, column=2, value=value)
        if label in ("Category Breakdown", "University", "Total Programs", "Total Facts"):
            cell_a.font = label_font
            cell_b.font = label_font
        else:
            cell_a.font = value_font
            cell_b.font = value_font

    # Save
    wb.save(output_path)
    print(f"\n  [DONE] Excel saved to: {output_path}")
    print(f"         {len(programs)} programs, {len(facts)} total facts")


def main():
    parser = argparse.ArgumentParser(description="Export extraction data to Excel")
    parser.add_argument(
        "--university",
        required=True,
        choices=list(UNIVERSITY_PATHS.keys()),
        help="University to export (asu, frankfurt, limerick)",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output Excel file path (default: <university>_extraction_data.xlsx)",
    )
    args = parser.parse_args()

    uni_path = UNIVERSITY_PATHS[args.university]
    output_path = Path(args.output) if args.output else Path(f"{args.university}_extraction_data.xlsx")

    print(f"\n  Exporting {args.university.upper()} data...")
    print(f"  Source: {uni_path}")

    facts = collect_all_facts(uni_path)

    if not facts:
        print("  [ERROR] No facts found. Check that extraction has completed.")
        sys.exit(1)

    write_excel(facts, output_path, args.university)


if __name__ == "__main__":
    main()
